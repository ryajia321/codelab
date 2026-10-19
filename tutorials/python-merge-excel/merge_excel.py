#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 汇总小助手 v1.0
把「演示目录/输入/」下的多个 .xlsx 文件合并成一张总表，
写到「演示目录/输出/汇总结果.xlsx」（附带"来源文件"列方便核对）。

安全策略：
  1. 只读「演示目录/输入/」顶层文件，不递归，不扫描电脑其他位置；
  2. 绝不修改、删除原始文件（只读方式打开）；
  3. 输出文件已存在时默认拒绝覆盖，提示用 --output 换名或加 --force 明确覆盖；
  4. 输出位置必须在演示目录内，且不能在输入目录里（避免下次运行重复汇总）；
  5. 列名不一致、空文件、损坏文件、重复表头：逐个文件明确报告并跳过，不静默错位；
  6. 来源符号链接、隐藏文件默认跳过；旧格式 .xls 明确提示不支持；
  7. 演示目录必须位于用户主目录内（realpath 解析后判断），拒绝根目录/主目录本身/符号链接入口。

依赖：pip install pandas openpyxl（只支持 .xlsx 格式）

用法：
  python3 merge_excel.py 演示目录                  # 输入=演示目录/输入，输出=演示目录/输出/汇总结果.xlsx
  python3 merge_excel.py 演示目录 --output 总表2.xlsx
  python3 merge_excel.py 演示目录 --force           # 允许覆盖已存在的输出文件

退出码：0 成功；1 运行失败（无文件可合并/输出已存在等）；2 用法或安全拒绝。
"""
import argparse
import os
import sys

import pandas as pd
from openpyxl import load_workbook


def real_home():
    return os.path.realpath(os.path.expanduser("~"))


def check_demo(demo):
    """返回 (真实路径, None) 或 (None, 错误信息)。"""
    if os.path.islink(demo):
        return None, "拒绝：入口不能是符号链接（当前：%s）" % demo
    demo = os.path.realpath(demo)
    if not os.path.isdir(demo):
        return None, "目录不存在：%s" % demo
    home = real_home()
    if demo == os.path.sep:
        return None, "拒绝：不能使用系统根目录"
    if demo == home:
        return None, "拒绝：不能直接使用用户主目录，请指定主目录下的演示目录"
    if not demo.startswith(home + os.sep):
        return None, "拒绝：只能使用用户主目录内的演示目录（当前：%s）" % demo
    return demo, None


def raw_headers(path):
    """用 openpyxl 直接读第一行表头（pandas 读进来会把重复列名静默改成"商品.1"，看不出重复）。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    finally:
        wb.close()
    return [str(v) if v is not None else "" for v in (row or [])]


def read_one(path):
    """读取单个 Excel。返回 (df, None) 或 (None, 跳过原因)。只读，不修改文件。"""
    try:
        headers = raw_headers(path)
    except Exception as e:  # 损坏文件、密码保护等都在这里兜住
        return None, "文件损坏或无法读取（%s），已跳过" % e
    if len(headers) != len(set(headers)):
        return None, "表头有重复的列名，已跳过（请先修正表头再汇总）"
    try:
        df = pd.read_excel(path, engine="openpyxl")
    except Exception as e:
        return None, "文件损坏或无法读取（%s），已跳过" % e
    if len(df) == 0:
        return None, "空文件（只有表头、没有数据行），已跳过"
    return df, None


def main():
    ap = argparse.ArgumentParser(description="汇总演示目录/输入/下的 Excel（默认不覆盖输出）")
    ap.add_argument("demo_dir", help="演示目录路径（其下应有 输入/ 文件夹）")
    ap.add_argument("--output", default="汇总结果.xlsx", help="输出文件名（默认 汇总结果.xlsx，放在演示目录/输出/下）")
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的输出文件")
    a = ap.parse_args()

    demo, err = check_demo(a.demo_dir)
    if err:
        print(err)
        return 2

    in_dir = os.path.join(demo, "输入")
    if not os.path.isdir(in_dir):
        print("输入文件夹不存在：%s\n请先在演示目录下建好 输入/ 文件夹并放入 Excel 文件。" % in_dir)
        return 2

    out_dir = os.path.join(demo, "输出")
    # --output 允许相对名（放进 输出/）或绝对路径（必须仍在演示目录内）
    out_path = a.output if os.path.isabs(a.output) else os.path.join(out_dir, a.output)
    out_path = os.path.realpath(out_path)
    if out_path != demo and not out_path.startswith(demo + os.sep):
        print("拒绝：输出文件必须在演示目录内（当前：%s）" % out_path)
        return 2
    if out_path == in_dir or out_path.startswith(in_dir + os.sep):
        print("拒绝：输出文件不能放在输入目录里，否则下次运行会把它当成输入重复汇总。")
        return 2
    if os.path.lexists(out_path) and not a.force:
        print("输出文件已存在，拒绝覆盖：%s" % out_path)
        print("请用 --output 新文件名 指定别的名字，或确认要覆盖时加 --force。")
        return 1

    try:
        names = sorted(os.listdir(in_dir))
    except OSError as e:
        print("读取输入文件夹失败：%s" % e)
        return 1

    frames, skipped, expected_cols = [], [], None
    total_rows = 0
    for name in names:
        src = os.path.join(in_dir, name)
        if os.path.islink(src):
            skipped.append((name, "符号链接，默认跳过（不跟随）"))
            continue
        if os.path.isdir(src):
            continue  # 不递归
        if name.startswith("."):
            skipped.append((name, "隐藏文件，跳过"))
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext == ".xls":
            skipped.append((name, "旧格式 .xls 不支持，请用 Excel 另存为 .xlsx 后再汇总"))
            continue
        if ext != ".xlsx":
            skipped.append((name, "不是 .xlsx 文件，跳过"))
            continue
        df, reason = read_one(src)
        if reason:
            skipped.append((name, reason))
            continue
        cols = [str(c) for c in df.columns]
        if expected_cols is None:
            expected_cols = cols
        elif cols != expected_cols:
            skipped.append((name, "列名与第一个文件不一致，已跳过\n      期望列：%s\n      实际列：%s" % (expected_cols, cols)))
            continue
        df = df.copy()
        df["来源文件"] = name
        frames.append(df)
        total_rows += len(df)
        print("已读取：%s（%d 行）" % (name, len(df)))

    if skipped:
        print("\n===== 跳过 %d 个文件 =====" % len(skipped))
        for name, reason in skipped:
            print("  %s ：%s" % (name, reason))

    if not frames:
        print("\n没有可合并的文件，未生成汇总。")
        return 1

    merged = pd.concat(frames, ignore_index=True)
    try:
        os.makedirs(out_dir, exist_ok=True)
        merged.to_excel(out_path, index=False, engine="openpyxl")
    except OSError as e:
        print("写入输出文件失败：%s" % e)
        return 1

    print("\n===== 汇总完成 =====")
    print("合并文件数：%d，数据总行数：%d" % (len(frames), total_rows))
    print("输出文件：%s" % out_path)
    print("提示：用 Excel 打开输出文件，核对总行数，并可用「来源文件」列追溯每一行来自哪个文件。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
