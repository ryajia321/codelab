#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 清理小助手 v1.0
清理一份存在脏数据的 Excel：去重标记、删空白行、去首尾空格、统一日期/金额/数量格式。

只用 openpyxl（不用 pandas 做类型推断——pandas 会改写原始单元格类型，
例如把重复列名改成"商品.1"、把文本数字转成数字，见 v0.4 方案 §1）。

三档模式：
  --check               只在终端显示完整问题报告，不创建任何文件（严格零写入）
  （默认）              修正单元格 + 删除完全空白行，不删除任何含数据的行
  --deduplicate-exact   显式删除整行完全相同的重复（保留首条）

支持范围（保守策略，先讲清楚再动手）：
  1. 只支持 .xlsx（--output 其它扩展名直接拒绝）；
  2. 不支持含公式的工作簿——检测到公式单元格直接拒绝，
     因为输出阶段无法可靠保留公式及其依赖关系，绝不静默把公式变成数值；
  3. 不支持带时间部分（时分秒非零）的日期时间——检测到直接拒绝，
     绝不静默截断时间；
  4. 金额/数量只按"先验证完整格式、再转换"的严格规则处理，
     格式不明的一律保留原文 + 警告。

安全策略：
  1. 原始文件只读，运行前后 checksum 不变；
  2. 重复判定基于 openpyxl 读取的【原始快照】（值+类型），发生在任何清洗之前，
     禁止先规范化再判定；默认只标记不删除；
  3. 演示目录/输入、演示目录/输出 本身是符号链接一律拒绝（无论指向何处）；
     输出目标本身是符号链接一律拒绝；输出必须位于 演示目录/输出/ 内；
     输出不得等于输入源文件（--force 也拒绝）；
  4. 输出已存在默认拒绝覆盖（原子 os.link，无检查-创建竞态）；
     --force 仅在路径安全检查通过后做原子 os.replace；
  5. Excel 与报告分别原子发布，不谎称整体事务；报告失败 = 部分成功状态（如实报告）；
  6. 失败时只删除自己创建的临时文件，绝不碰用户原有文件；
  7. 演示目录必须位于用户主目录内（realpath 解析），拒绝根目录/主目录本身/符号链接入口。

威胁边界（诚实说明，不夸大）：
  realpath 解析 + islink 检查只能反映"检查时刻"的状态，不能消除检查与使用
  之间的竞态（TOCTOU）——有演示目录写权限的攻击者理论上可在检查后替换
  符号链接。本脚本面向用户自己的演示目录，上述检查用于防止误操作和常规
  风险，不声称能防御有本地写权限的恶意攻击者。

测试注入点（供自动化测试覆盖"报告发布失败"的部分成功路径，教程可忽略）：
  环境变量 CLEAN_EXCEL_INJECT_REPORT_FAIL=1 时，在 Excel 发布成功后模拟报告发布失败。

依赖：pip install openpyxl（只支持 .xlsx）
用法：
  python3 clean_excel.py 演示目录 --check
  python3 clean_excel.py 演示目录
  python3 clean_excel.py 演示目录 --deduplicate-exact
  python3 clean_excel.py 演示目录 --output 干净订单.xlsx [--force]
退出码：0 成功；1 运行失败（含部分成功）；2 用法或安全拒绝。
"""
import argparse
import math
import os
import re
import sys
import tempfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook

REQUIRED_COLS = ["订单号", "客户", "商品", "数量", "金额", "下单日期"]
INPUT_NAME = "脏订单.xlsx"
OUTPUT_NAME = "清理后.xlsx"
REPORT_NAME = "清理报告.txt"

DATE_FMTS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日")
QTY_UNITS = ("件", "个", "只", "瓶", "箱", "台", "套", "包")
AMBIGUOUS_DOT = re.compile(r"^\d{1,3}\.\d{3}$")  # 如 1.234：欧陆千分位歧义，不解释
NONFINITE_RE = re.compile(r"^[+-]?\s*(nan|inf|infinity)\s*$", re.IGNORECASE)
# 金额允许的完整格式（先整体验证，再转换）：
#   可选 1 个货币符号开头（¥/￥/$），或 "元" 结尾；
#   数字部分：纯数字（可选小数点）或规范千分位（每组恰 3 位）。
# 不在其中的（如 1,2、美元450、-450）一律保留原文 + 警告。
AMOUNT_RE = re.compile(r"^(?:[¥￥$]\s*)?(\d+(?:\.\d+)?|\d{1,3}(?:,\d{3})+(?:\.\d+)?)(?:\s*元)?$")
# 数量允许的完整格式：可选负号 + 整数 + 可选 1 个单位（单位必须在末尾）。
QTY_RE = re.compile(r"^(-?\d+)\s*(件|个|只|瓶|箱|台|套|包)?$")
QTY_DECIMAL_RE = re.compile(r"^-?\d+\.\d+$")


def real_home():
    return os.path.realpath(os.path.expanduser("~"))


def check_demo(demo):
    """演示目录入口检查。返回 (realpath, None) 或 (None, 错误信息)。

    威胁边界：realpath + islink 检查只反映检查时刻的状态，不能消除
    检查与使用之间的竞态（TOCTOU）。面向用户自己的演示目录，
    用于防止误操作和常规风险，不防御有本地写权限的恶意攻击者。
    """
    if os.path.islink(demo):
        return None, "拒绝：入口不能是符号链接（当前：%s）" % demo
    demo = os.path.realpath(demo)
    if not os.path.isdir(demo):
        return None, "目录不存在：%s" % demo
    home = real_home()
    if demo == os.path.sep:
        return None, "拒绝：不能使用系统根目录"
    if demo == home:
        return None, "拒绝：不能直接使用用户主目录，请指定主目录下的演示目录"
    if not demo.startswith(home + os.sep):
        return None, "拒绝：只能使用用户主目录内的演示目录（当前：%s）" % demo
    # 演示目录/输入 与 演示目录/输出 本身不得是符号链接（无论指向何处）
    for sub in ("输入", "输出"):
        if os.path.islink(os.path.join(demo, sub)):
            return None, "拒绝：演示目录/%s 本身是符号链接（无论指向何处都不允许）" % sub
    return demo, None


# ---------- 原始快照（openpyxl 两遍读取，不信任 pandas 类型推断） ----------

def scalar_key(val):
    """单个值的归一键。

    诚实口径（只承诺实际做到的）：
    - 用的是 openpyxl 返回的 Python 值（不是 Excel 底层存储字节）；
    - datetime 保留完整时分秒（不同时间 ≠ 相同）；
    - 数字 int/float 按 Decimal(str(v)).normalize() 比较，
      即整数 10 与浮点 10.0 视为相同；文本 "10" 与数字 10 类型不同，不视为相同；
    - 非有限数值（NaN/Infinity）单独标记，不参与相等比较。
    """
    if val is None:
        return ("empty",)
    if isinstance(val, bool):
        return ("bool", str(val))
    if isinstance(val, datetime):
        return ("date", val.isoformat())  # 保留时分秒
    if isinstance(val, date):
        return ("date", val.isoformat())
    if isinstance(val, (int, float)):
        if isinstance(val, float) and not math.isfinite(val):
            return ("nonfinite",)
        return ("num", str(Decimal(str(val)).normalize()))
    return ("text", str(val))


def cell_key(val, dtype, formula):
    """单个单元格的原始快照键。val 来自 data_only=True，dtype/formula 来自 data_only=False。

    dtype 是 openpyxl 报告的单元格类型（f=公式，e=错误值，s=文本，n=数字，d=日期）。
    """
    if dtype == "f":
        if val is None:
            return ("formula-missing",)
        return ("formula", str(formula), scalar_key(val))
    if dtype == "e":
        return ("error",)  # Excel 错误值（#DIV/0! 等）：不参与相等比较
    return scalar_key(val)


def read_rows(path):
    """返回 (header, rows)。rows 每项：{'excel_row': n, 'cells': [(val, dtype, formula), ...], 'keys': tuple}。"""
    wb_v = load_workbook(path, read_only=True, data_only=True)
    wb_r = load_workbook(path, read_only=True, data_only=False)
    try:
        wv, wr = wb_v.active, wb_r.active
        header = [c.value for c in next(wv.iter_rows(min_row=1, max_row=1))]
        rows = []
        for rn, (rv, rr) in enumerate(zip(wv.iter_rows(min_row=2), wr.iter_rows(min_row=2)), start=2):
            cells, keys = [], []
            for cv, cr in zip(rv, rr):
                formula = cr.value if cr.data_type == "f" else None
                cells.append((cv.value, cr.data_type, formula))
                keys.append(cell_key(cv.value, cr.data_type, formula))
            rows.append({"excel_row": rn, "cells": cells, "keys": tuple(keys)})
    finally:
        wb_v.close()
        wb_r.close()
    return header, rows


def is_blank(cells):
    """空白行：所有格为空或纯空白文本（独立于重复判定）。

    公式单元格（dtype == 'f'）即使缓存值为空也视为有内容——
    含公式的工作簿会被上游直接拒绝，此处是纵深防御，绝不把公式行当空白行删掉。
    """
    for val, dtype, _ in cells:
        if dtype == "f":
            return False
        if val is None:
            continue
        if isinstance(val, str) and val.strip() == "":
            continue
        return False
    return True


def find_exact_groups(rows):
    """返回 (groups, skipped)。

    groups: {keys: [excel_row,...]}，只含 >= 2 行的组（可安全判定的完全重复）。
    skipped: [(excel_row, 原因)]——含公式无缓存 / Excel 错误值 / 非有限数值 的行，
    不参与重复判定（宁可漏判，不可误判），调用方应给出明确警告。
    """
    bucket, skipped = {}, []
    for r in rows:
        reason = None
        for k in r["keys"]:
            if k[0] == "formula-missing":
                reason = "公式无缓存值，视为缺失"
                break
            if k[0] == "error":
                reason = "含 Excel 错误值（如 #DIV/0!）"
                break
            if k[0] == "nonfinite":
                reason = "含非有限数值（NaN/Infinity）"
                break
        if reason:
            skipped.append((r["excel_row"], reason))
            continue
        bucket.setdefault(r["keys"], []).append(r["excel_row"])
    groups = {k: v for k, v in bucket.items() if len(v) > 1}
    return groups, skipped


def find_suspects(rows, oid_idx):
    """订单号相同但快照不同 → 疑似重复，始终保留+警告。返回 [(oid, [excel_row,...])]。"""
    bucket = {}
    for r in rows:
        val = r["cells"][oid_idx][0]
        if val is None:
            continue
        bucket.setdefault(str(val), []).append(r)
    out = []
    for oid, rs in bucket.items():
        if len(rs) > 1 and len({r["keys"] for r in rs}) > 1:
            out.append((oid, [r["excel_row"] for r in rs]))
    return out


# ---------- 清洗（只作用于保留行；返回 (新值, 是否实际改写, 警告)） ----------

def clean_text(val):
    if not isinstance(val, str):
        return val, False, None
    s = val.strip()
    return s, (s != val), None


def clean_date(val):
    # 注意：带时间部分（时分秒非零）的 datetime 已在上游被拒绝，
    # 到这里的 datetime 一定是 00:00:00，只统一写法，不计修正。
    if isinstance(val, datetime):
        return val.date().isoformat(), False, None
    if isinstance(val, date):
        return val.isoformat(), False, None
    s = str(val).strip()
    for f in DATE_FMTS:
        try:
            iso = datetime.strptime(s, f).date().isoformat()
            return iso, (iso != s), None  # 与原文相同不计修正
        except ValueError:
            pass
    return val, False, "日期无法解析（%s），保留原文" % val


def clean_amount(val):
    """金额：先验证完整格式，再转换。格式不明的一律保留原文 + 警告。

    允许：可选 1 个货币符号（¥/￥/$）开头或"元"结尾 + 纯数字/规范千分位（可带小数）。
    拒绝：1,2（千分位不规范）、1.234（歧义）、NaN/Infinity、美元450、-450 等。
    """
    if isinstance(val, bool):
        return val, False, "金额是布尔值，保留原文"
    if isinstance(val, (int, float)):
        if isinstance(val, float) and not math.isfinite(val):
            return val, False, "金额是非有限数值，保留原文"
        return val, False, None
    s = str(val).strip()
    if not s:
        return val, False, "金额为空，保留原文"
    if NONFINITE_RE.match(s):
        return val, False, "金额是非有限数值（%s），保留原文" % val
    if AMBIGUOUS_DOT.match(s) or s.count(".") > 1:
        return val, False, "金额格式歧义（%s），不擅自解释，保留原文" % val
    m = AMOUNT_RE.match(s)
    if not m:
        return val, False, "金额格式无法确认（%s），保留原文" % val
    try:
        d = Decimal(m.group(1).replace(",", ""))
    except InvalidOperation:
        return val, False, "金额无法转换为数字（%s），保留原文" % val
    if not d.is_finite():
        return val, False, "金额是非有限数值，保留原文"
    v = int(d) if d == d.to_integral_value() else float(d)
    return v, True, None


def clean_qty(val):
    """数量：先验证完整格式，再转换。格式不明的一律保留原文 + 警告。

    允许：整数（可选负号）+ 可选 1 个单位（件/个/只/瓶/箱/台/套/包，须在末尾）。
    拒绝：1件2（单位不在末尾）、3.5（不是整数）、NaN/Infinity 等。
    """
    if isinstance(val, bool):
        return val, False, "数量是布尔值，保留原文"
    if isinstance(val, (int, float)):
        if isinstance(val, float) and not math.isfinite(val):
            return val, False, "数量是非有限数值，保留原文"
        return val, False, None
    s = str(val).strip()
    if not s:
        return val, False, "数量为空，保留原文"
    if NONFINITE_RE.match(s):
        return val, False, "数量是非有限数值（%s），保留原文" % val
    if QTY_DECIMAL_RE.match(s):
        return val, False, "数量不是整数（%s），保留原文" % val
    m = QTY_RE.match(s)
    if not m:
        return val, False, "数量格式无法确认（%s），保留原文" % val
    return int(m.group(1)), True, None


def clean_row(cells, col_idx):
    """清洗一行保留行。返回 (new_values, fixed_count, warnings[(列名, 信息)])。"""
    out, fixed, warns = [], 0, []
    for i, (val, _, _) in enumerate(cells):
        col = col_idx["names"][i]
        if col in ("客户", "商品"):
            nv, ch, w = clean_text(val)
        elif col == "下单日期":
            nv, ch, w = clean_date(val)
        elif col == "金额":
            nv, ch, w = clean_amount(val)
        elif col == "数量":
            nv, ch, w = clean_qty(val)
        else:  # 订单号等：保持原样，不碰身份字段
            nv, ch, w = val, False, None
        out.append(nv)
        if ch:
            fixed += 1
        if w:
            warns.append((col, w))
    return out, fixed, warns


# ---------- 路径安全 ----------

def resolve_in_dir(demo, sub, arg, must_exist=False):
    """把 arg 解析到 演示目录/sub/ 内。返回 (real_path, None) 或 (None, (msg, is_safety))。

    威胁边界：realpath 解析只能反映检查时刻的状态，不能消除检查与使用
    之间的竞态（TOCTOU）。配合 check_demo 对输入/输出目录本身的 islink
    检查，用于防止误操作和常规风险，不防御有本地写权限的恶意攻击者。
    """
    base = os.path.join(demo, sub)
    real_base = os.path.realpath(base)
    if real_base != demo and not real_base.startswith(demo + os.sep):
        return None, ("拒绝：%s 目录解析后超出演示目录范围" % sub, True)
    # 输出目录不得解析进输入目录（避免输出落进输入，下次被当成输入）
    if sub == "输出":
        real_in = os.path.realpath(os.path.join(demo, "输入"))
        if real_base == real_in or real_base.startswith(real_in + os.sep):
            return None, ("拒绝：输出目录不能指向输入目录内", True)
    raw = arg if os.path.isabs(arg) else os.path.join(base, arg)
    if os.path.islink(raw):
        return None, ("拒绝：目标本身是符号链接（%s），无论指向何处都不允许" % arg, True)
    final = os.path.realpath(raw)
    if final != real_base and not final.startswith(real_base + os.sep):
        return None, ("拒绝：目标必须在演示目录/%s/ 内（当前：%s）" % (sub, arg), True)
    if os.path.isdir(final):
        return None, ("拒绝：输出目标是目录，无法写入文件：%s" % arg, True)
    if must_exist and not os.path.isfile(final):
        return None, ("输入文件不存在：%s" % final, True)
    return final, None


# ---------- 原子发布 ----------

def publish_file(tmp, final, force):
    """原子发布单个文件。返回 (True, None) 或 (False, 原因)。"""
    try:
        if force:
            os.replace(tmp, final)
        else:
            os.link(tmp, final)  # 目标已存在则抛 FileExistsError：无检查-创建竞态
    except FileExistsError:
        return False, "目标已存在，拒绝覆盖：%s" % final
    except OSError as e:
        return False, "发布失败（%s）：%s" % (final, e)
    return True, None


# ---------- 报告 ----------

def row_text(cells):
    parts = []
    for v, _, _ in cells:
        if v is None:
            parts.append("")
        elif isinstance(v, datetime):
            parts.append(v.date().isoformat())
        elif isinstance(v, date):
            parts.append(v.isoformat())
        else:
            parts.append(str(v))
    return " | ".join(parts)


def build_report(ctx):
    L = []
    L.append("===== Excel 清理报告 =====")
    L.append("输入文件：%s" % ctx["input"])
    L.append("模式：%s" % ctx["mode"])
    L.append("原始数据行数：%d" % ctx["n_raw"])
    L.append("自动修正单元格数：%d" % ctx["n_fixed"])
    L.append("删除行数：%d" % ctx["n_deleted"])
    L.append("警告数：%d" % ctx["n_warn"])
    if ctx["groups"]:
        L.append("")
        L.append("--- 重复组（整行原始值+类型完全相同） ---")
        for gi, (key, rows) in enumerate(ctx["groups"].items(), start=1):
            L.append("重复组 G%d（共 %d 行）：" % (gi, len(rows)))
            for j, rn in enumerate(rows):
                tag = "保留（首条）" if j == 0 else ("已删除" if ctx["dedup"] else "标记未删除")
                L.append("  - Excel第%d行：%s → %s" % (rn, ctx["row_text"][rn], tag))
    if ctx["suspects"]:
        L.append("")
        L.append("--- 疑似重复（订单号相同、内容不同，始终保留） ---")
        for oid, rns in ctx["suspects"]:
            L.append("  - 订单号 %s：Excel第%s行 → 保留全部，请人工确认" % (oid, "、第".join(map(str, rns))))
    if ctx["warnings"]:
        L.append("")
        L.append("--- 警告明细 ---")
        for rn, col, msg in ctx["warnings"]:
            L.append("  - Excel第%d行【%s】：%s" % (rn, col, msg))
    if ctx["deleted"]:
        L.append("")
        L.append("--- 删除明细 ---")
        for rn, reason in ctx["deleted"]:
            L.append("  - Excel第%d行：%s" % (rn, reason))
    if not ctx["groups"] and not ctx["suspects"] and not ctx["warnings"] and not ctx["deleted"]:
        L.append("未发现问题，数据干净。")
    return "\n".join(L) + "\n"


# ---------- 主流程 ----------

def main():
    ap = argparse.ArgumentParser(description="清理 Excel 脏数据（默认不删除业务行）")
    ap.add_argument("demo_dir", help="演示目录路径（其下应有 输入/ 文件夹）")
    ap.add_argument("--input", default=INPUT_NAME, help="输入文件名（默认 %s，位于演示目录/输入/ 下）" % INPUT_NAME)
    ap.add_argument("--output", default=OUTPUT_NAME, help="输出文件名（默认 %s，位于演示目录/输出/ 下，仅支持 .xlsx）" % OUTPUT_NAME)
    ap.add_argument("--check", action="store_true", help="只显示问题报告，不创建任何文件")
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的输出文件（仅路径安全时）")
    ap.add_argument("--deduplicate-exact", action="store_true", help="显式删除整行完全相同的重复（保留首条）")
    a = ap.parse_args()

    # --output 只支持 .xlsx（其它扩展名直接拒绝，避免报告名派生出误导性文件名）
    if not a.output.lower().endswith(".xlsx"):
        print("拒绝：--output 只支持 .xlsx 文件（当前：%s）。" % a.output)
        print("本工具只读写 .xlsx 格式（Excel/WPS 默认格式），请改用 .xlsx 后缀。")
        return 2

    demo, err = check_demo(a.demo_dir)
    if err:
        print(err)
        return 2

    in_path, err = resolve_in_dir(demo, "输入", a.input, must_exist=True)
    if err:
        print(err[0])
        return 2 if err[1] else 1
    # 注：resolve_in_dir 已拒绝符号链接输入；realpath 解析后仍在 输入/ 内才放行

    try:
        header, rows = read_rows(in_path)
    except Exception as e:
        print("读取输入文件失败：%s" % e)
        return 1

    missing = [c for c in REQUIRED_COLS if c not in header]
    if missing:
        print("缺少必要列：%s\n期望列：%s" % ("、".join(missing), "、".join(REQUIRED_COLS)))
        return 1
    if not rows:
        print("表格没有数据行，未执行任何操作。")
        return 1
    col_idx = {"names": header, "map": {name: header.index(name) for name in header}}
    oid_idx = col_idx["map"]["订单号"]

    # ---- 支持范围检查（保守策略）：在任何判定/写入之前拒绝 ----
    # 含公式的工作簿：输出阶段无法可靠保留公式及其依赖关系，绝不静默丢弃公式。
    formula_rows = sorted({r["excel_row"] for r in rows
                           if any(dtype == "f" for _, dtype, _ in r["cells"])})
    if formula_rows:
        print("拒绝：检测到公式单元格（Excel第%s行）。" % "、第".join(map(str, formula_rows)))
        print("当前版本无法可靠地保留公式及其计算关系，为避免静默丢弃公式，已拒绝自动清理。")
        print("原始文件未改动，也没有创建任何输出文件。")
        print("如需处理含公式的表，请先在 Excel/WPS 中把公式计算结果「粘贴为数值」后再试。")
        return 2
    # 带时间部分的日期时间：绝不静默截断时分秒。
    timed = []
    for r in rows:
        for i, (val, _, _) in enumerate(r["cells"]):
            if isinstance(val, time) or (
                    isinstance(val, datetime)
                    and (val.hour, val.minute, val.second, val.microsecond) != (0, 0, 0, 0)):
                timed.append((r["excel_row"], header[i]))
    if timed:
        shown = "、".join("Excel第%d行【%s】" % t for t in timed[:5])
        if len(timed) > 5:
            shown += "（等共 %d 处）" % len(timed)
        print("拒绝：检测到带时间部分的日期时间（%s）。" % shown)
        print("当前版本不会静默截断时间，已拒绝自动清理。原始文件未改动，也没有创建任何输出文件。")
        return 2

    # 判定（清洗之前，基于原始快照）
    groups, skipped = find_exact_groups(rows)
    suspects = find_suspects(rows, oid_idx)

    # 删除集合
    to_delete = {}  # excel_row -> 原因
    for r in rows:
        if is_blank(r["cells"]):
            to_delete[r["excel_row"]] = "完全空白行，已删除"
    if a.deduplicate_exact:
        for key, rns in groups.items():
            for rn in rns[1:]:
                to_delete[rn] = "整行完全相同（显式去重），已删除"

    # 清洗保留行
    out_rows, n_fixed, warnings, row_texts = [], 0, {}, {}
    for r in rows:
        rn = r["excel_row"]
        row_texts[rn] = row_text(r["cells"])
        if rn in to_delete:
            continue
        new_vals, fixed, warns = clean_row(r["cells"], col_idx)
        out_rows.append(new_vals)
        n_fixed += fixed
        for col, msg in warns:
            warnings.setdefault(rn, []).append((col, msg))
    for rn, reason in skipped:
        warnings.setdefault(rn, []).append(("整行", "%s，不参与重复判定" % reason))

    # 警告计数：重复组标记 + 疑似组 + 单元格警告（去重模式下已删除的组不计警告）
    n_warn = 0
    if not a.deduplicate_exact:
        n_warn += len(groups)
    n_warn += len(suspects)
    for rn, lst in warnings.items():
        n_warn += len(lst)

    flat_warnings = []
    for gi, (key, rns) in enumerate(groups.items(), start=1):
        if not a.deduplicate_exact:
            flat_warnings.append((rns[0], "重复组", "G%d 共 %d 行完全相同，已标记未删除" % (gi, len(rns))))
    for oid, rns in suspects:
        flat_warnings.append((rns[0], "订单号", "疑似重复（订单号 %s），已保留全部" % oid))
    for rn, lst in warnings.items():
        for col, msg in lst:
            flat_warnings.append((rn, col, msg))
    flat_warnings.sort()

    deleted_list = sorted((rn, to_delete[rn]) for rn in to_delete)
    ctx = {
        "input": in_path,
        "mode": "--check 预览" if a.check else ("显式去重" if a.deduplicate_exact else "默认执行"),
        "n_raw": len(rows),
        "n_fixed": n_fixed,
        "n_deleted": len(to_delete),
        "n_warn": n_warn,
        "groups": groups,
        "suspects": suspects,
        "warnings": flat_warnings,
        "deleted": deleted_list,
        "dedup": a.deduplicate_exact,
        "row_text": row_texts,
    }
    report = build_report(ctx)

    if a.check:
        # 严格零写入：到此为止未创建任何文件系统对象，直接打印
        print(report)
        return 0

    # 写前预验证（两个目标，全部通过才发布任一；报告名随输出名派生）
    xlsx_final, err = resolve_in_dir(demo, "输出", a.output)
    if err:
        print(err[0])
        return 2 if err[1] else 1
    report_name = os.path.splitext(os.path.basename(a.output))[0] + "_报告.txt"
    report_final, err = resolve_in_dir(demo, "输出", report_name)
    if err:
        print(err[0])
        return 2 if err[1] else 1
    if not a.force:
        if os.path.lexists(xlsx_final):
            print("输出文件已存在，拒绝覆盖：%s\n请用 --output 换名，或确认要覆盖时加 --force。" % xlsx_final)
            return 1
        if os.path.lexists(report_final):
            print("报告文件已存在，拒绝覆盖：%s\n请确认要覆盖时加 --force。" % report_final)
            return 1

    real_out_dir = os.path.realpath(os.path.join(demo, "输出"))
    os.makedirs(real_out_dir, exist_ok=True)

    temps = []
    try:
        # 准备临时文件（mkstemp：O_CREAT|O_EXCL + 随机名；xlsx 临时文件保留 .xlsx 后缀以便校验）
        fd, xlsx_tmp = tempfile.mkstemp(dir=real_out_dir, prefix=".clean-xlsx-", suffix=".xlsx")
        temps.append(xlsx_tmp)
        fd2, report_tmp = tempfile.mkstemp(dir=real_out_dir, prefix=".clean-report-", suffix=".txt")
        temps.append(report_tmp)

        wb = Workbook()
        ws = wb.active
        ws.title = "清理后"
        ws.append(header)
        for vals in out_rows:
            ws.append(vals)
        wb.save(xlsx_tmp)
        os.close(fd)
        with os.fdopen(fd2, "w", encoding="utf-8") as f:
            f.write(report)

        # 校验临时文件
        try:
            wv = load_workbook(xlsx_tmp, read_only=True)
            wv.close()
        except Exception as e:
            print("生成的 Excel 校验失败，中止：%s" % e)
            return 1
        if os.path.getsize(report_tmp) == 0:
            print("生成的报告为空，中止。")
            return 1

        # 发布 Excel
        ok, msg = publish_file(xlsx_tmp, xlsx_final, a.force)
        if not ok:
            print(msg)
            return 1
        temps.remove(xlsx_tmp)
        try:
            os.unlink(xlsx_tmp)  # os.link 后需手动删；os.replace 改名后已不存在
        except FileNotFoundError:
            pass

        # 测试注入点：模拟报告发布失败（覆盖部分成功路径）
        if os.environ.get("CLEAN_EXCEL_INJECT_REPORT_FAIL") == "1":
            ok, msg = False, "注入的报告发布失败（测试用 CLEAN_EXCEL_INJECT_REPORT_FAIL=1）"
        else:
            ok, msg = publish_file(report_tmp, report_final, a.force)
        if not ok:
            print("===== 部分成功 =====")
            print("Excel 已生成（本次新输出）：%s" % xlsx_final)
            print("报告写入失败：%s" % msg)
            print("报告缺失不影响 Excel 数据；可用 --force 重跑一次补齐报告。")
            print("注意：这不是整体成功，已发布的 Excel 为本次有效输出，予以保留（不回滚）。")
            return 1
        temps.remove(report_tmp)
        try:
            os.unlink(report_tmp)
        except FileNotFoundError:
            pass

        print(report)
        print("输出文件：%s" % xlsx_final)
        print("报告文件：%s" % report_final)
        return 0
    finally:
        for t in temps:  # 只删除本次创建的临时文件
            try:
                os.unlink(t)
            except OSError:
                pass


if __name__ == "__main__":
    sys.exit(main())
