#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成 12 行脏数据示例（订单表），供《用 Python 清理 Excel 数据》教程演示用。
全部虚构数据，不含任何真实信息。

关键：用 openpyxl 直接写单元格，精确控制每个单元格的存储类型——
  第 2 行下单日期写成【真实日期单元格】，第 3 行写成【文本】"2026/10/02"，
  第 5 行与第 2 行逐格同值同类型（真重复），订单号 "00012" 写成文本（保前导零）。

安全：默认不覆盖已存在的目标文件（防止盖掉用户的真实文件）；
  目标本身是符号链接一律拒绝。只有加 --force 才允许覆盖。

用法：python3 make_dirty_samples.py 演示目录 [--force]
将在「演示目录/输入/」下生成 脏订单.xlsx（12 行数据 + 1 行表头）。
"""
import argparse
import os
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import numbers

HEADER = ["订单号", "客户", "商品", "数量", "金额", "下单日期"]

# 每行：(订单号, 客户, 商品, 数量, 金额, 下单日期, 日期是否写成真实日期单元格)
# 数量/金额中带引号的字符串表示【文本】类型，不带引号的数字表示【数字】类型
ROWS = [
    ("A001", "张三", "苹果", 10, 300, datetime(2026, 10, 1), True),    # 行2：干净，真实日期格
    ("A002", "李四", "香蕉", 20, 400, "2026/10/02", False),            # 行3：日期文本
    ("A003", "王五", "橙子", 15, "¥450", "2026-10-03", False),          # 行4：金额文本
    ("A001", "张三", "苹果", 10, 300, datetime(2026, 10, 1), True),    # 行5：真重复（逐格同值同类型）
    ("A004", "赵六", "梨", "3件", 180, "2026-10-04", False),            # 行6：数量文本
    (None, None, None, None, None, None, False),                       # 行7：完全空白行
    ("A005", "钱七 ", "苹果", 10, 300, "2026-10-05", False),            # 行8：客户首尾空格
    ("A006", "孙八", "香蕉", 5, "待定", "2026-10-06", False),           # 行9：金额无法转换
    ("A007", "周九", "橙子", 8, 240, "10/07", False),                  # 行10：日期缺少年份
    ("A008", "吴十", "梨", 12, 360, "2026-10-08", False),              # 行11：干净
    ("A002", "李四", "香蕉", 25, 500, "2026-10-02", False),            # 行12：疑似重复（同订单号，内容不同）
    ("00012", "郑十一", "葡萄", 6, 210, "2026-10-09", False),          # 行13：订单号前导零（文本）
]


def main():
    ap = argparse.ArgumentParser(description="生成 12 行脏数据示例（默认不覆盖已存在文件）")
    ap.add_argument("demo_dir", nargs="?", default=os.path.expanduser("~/excel-clean-demo"),
                    help="演示目录路径（默认 ~/excel-clean-demo）")
    ap.add_argument("--force", action="store_true", help="允许覆盖已存在的 脏订单.xlsx")
    a = ap.parse_args()

    demo = os.path.realpath(a.demo_dir)
    home = os.path.realpath(os.path.expanduser("~"))
    if demo == home or not demo.startswith(home + os.sep):
        print("拒绝：请指定主目录下的演示目录。")
        return 2
    in_dir = os.path.join(demo, "输入")
    os.makedirs(in_dir, exist_ok=True)

    path = os.path.join(in_dir, "脏订单.xlsx")
    # 目标本身是符号链接：无论指向何处，一律拒绝（防止通过链接覆盖用户真实文件）
    if os.path.islink(path):
        print("拒绝：目标本身是符号链接，不覆盖：%s" % path)
        return 2
    # 默认不覆盖已存在文件：样例生成器绝不能悄悄盖掉用户的真实文件
    if os.path.lexists(path) and not a.force:
        print("拒绝：目标文件已存在，不覆盖：%s" % path)
        print("如需重新生成示例，请确认后加 --force。")
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(HEADER)
    for oid, cust, prod, qty, amt, dt, real_date in ROWS:
        row = [oid, cust, prod, qty, amt, dt]
        ws.append(row)
        if real_date:
            # 把下单日期格设为真实日期类型（openpyxl 按 datetime 自动识别，此处显式确认）
            cell = ws.cell(row=ws.max_row, column=6)
            cell.number_format = numbers.FORMAT_DATE_YYYYMMDD2
    # 订单号列强制文本格式（保前导零 "00012"）
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "@"
    wb.save(path)
    print("已生成 脏订单.xlsx（12 行数据，全部虚构）→ %s" % path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
