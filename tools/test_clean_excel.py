#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""clean_excel.py 自动化测试（41 项）。
- 数据：程序生成的虚构数据；Excel 类型测试用 openpyxl 实际生成真 .xlsx 文件，
  不用 Python 列表模拟 Excel 类型。
- 目录：tempfile.TemporaryDirectory() 独立临时目录（主目录内），只删除本次创建的目录。
- 哨兵：放在本次测试专用的临时目录内，不再使用用户主目录下的固定文件名。
- 用法：python3 tools/test_clean_excel.py（在 codelab-site 目录下执行）"""
import glob
import hashlib
import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(ROOT, "tutorials", "python-clean-excel", "clean_excel.py")
SAMPLE = os.path.join(ROOT, "tutorials", "python-clean-excel", "make_dirty_samples.py")
HOME = os.path.expanduser("~")
HEADER = ["订单号", "客户", "商品", "数量", "金额", "下单日期"]
PASS, FAIL = 0, 0

def ok(n): globals()["PASS"] += 1; print("PASS", n)
def bad(n, w): globals()["FAIL"] += 1; print("FAIL", n, "--", w)

def run(demo, *args, env=None):
    e = dict(os.environ)
    if env: e.update(env)
    p = subprocess.run([sys.executable, SCRIPT, demo, *args], capture_output=True, text=True, env=e)
    return p.returncode, p.stdout + p.stderr

def run_gen(demo, *args):
    p = subprocess.run([sys.executable, SAMPLE, demo, *args], capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr

def md5(p): return hashlib.md5(open(p, "rb").read()).hexdigest()

def stats(out):
    d = {}
    for line in out.splitlines():
        for k in ("原始数据行数", "自动修正单元格数", "删除行数", "警告数"):
            if line.startswith(k + "："):
                d[k] = int(line.split("：")[1])
    return d

def tree(demo):
    files, dirs = {}, []
    for root, ds, fs in os.walk(demo):
        for dd in ds: dirs.append(os.path.relpath(os.path.join(root, dd), demo))
        for f in fs:
            rel = os.path.relpath(os.path.join(root, f), demo)
            files[rel] = md5(os.path.join(root, f))
    return sorted(files), files, sorted(dirs)

def write_xlsx(path, rows):
    """rows: 每格为普通值，或 (值, 类型) 元组，类型 in text/formula。"""
    wb = Workbook(); ws = wb.active
    ws.append(HEADER)
    for row in rows:
        vals = [v[0] if isinstance(v, tuple) else v for v in row]
        ws.append(vals)
        for ci, v in enumerate(row, start=1):
            if isinstance(v, tuple):
                val, typ = v
                c = ws.cell(row=ws.max_row, column=ci)
                if typ == "text":
                    c.value = str(val); c.data_type = "s"
                elif typ == "formula":
                    c.value = val  # "=..." openpyxl 自动识别为公式（无缓存值）
    wb.save(path)

def load_mod():
    spec = importlib.util.spec_from_file_location("clean_excel", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

def main():
    global PASS, FAIL
    # 哨兵：本次测试专用的临时目录内（独占创建），绝不在用户主目录用固定文件名
    sen_dir = tempfile.mkdtemp(prefix="codelab-test-sentinel-", dir=HOME)
    sentinel = os.path.join(sen_dir, "sentinel.txt")
    with open(sentinel, "w", encoding="utf-8") as f:
        f.write("sentinel-ok")
    try:
        with tempfile.TemporaryDirectory(prefix="codelab-clean-test-", dir=HOME) as BASE:
            def fresh(name):
                d = os.path.join(BASE, name)
                os.makedirs(os.path.join(d, "输入"))
                return d
            def gen(d):
                rc, out = run_gen(d)
                assert rc == 0, out
            def inp(d): return os.path.join(d, "输入", "脏订单.xlsx")

            # 1. 12 行样例全量清理（默认模式）
            d = fresh("t1"); gen(d)
            h0 = md5(inp(d))
            rc, out = run(d)
            st = stats(out)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            if rc == 0 and st == {"原始数据行数": 12, "自动修正单元格数": 4, "删除行数": 1, "警告数": 4} \
                    and ws.max_row == 12 and md5(inp(d)) == h0:
                ok("1.全量清理：12行/修正4/删除1/警告4，原始文件未动")
            else: bad("1.全量清理", "rc=%d st=%s rows=%d" % (rc, st, ws.max_row))

            # 2. 真重复（行2/行5）：默认标记不删
            d = fresh("t2"); gen(d)
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            oids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
            if "重复组 G1" in out and oids.count("A001") == 2 and "标记未删除" in out:
                ok("2.真重复：默认标记不删，报告列出行号")
            else: bad("2.真重复", out[:200])

            # 2x. 显式去重：删除行5
            rc, out = run(d, "--deduplicate-exact", "--output", "去重.xlsx", "--force")
            st = stats(out)
            ws = load_workbook(os.path.join(d, "输出", "去重.xlsx")).active
            oids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
            if st["删除行数"] == 2 and ws.max_row == 11 and oids.count("A001") == 1 and "已删除" in out:
                ok("2x.显式去重：删除行5，报告标注保留/删除")
            else: bad("2x.显式去重", "st=%s" % st)

            # 3. 数字 10 vs 文本 "10"：真 Excel 文件，不判重复
            d = fresh("t3")
            write_xlsx(inp(d), [
                ["T01", "甲", "X", 10, 100, "2026-01-01"],
                ["T02", "甲", "X", ("10", "text"), 100, "2026-01-01"],
            ])
            rc, out = run(d)
            if rc == 0 and "重复组" not in out:
                ok("3.数字10 vs 文本'10'：真文件验证，不判重复")
            else: bad("3.类型判定", out[:200])

            # 4. 公式分组逻辑（内部函数单元测试：同公式字符串+同缓存值才进组）
            mod = load_mod()
            k1 = (mod.cell_key(5, "f", "=A1+B1"),)
            k2 = (mod.cell_key(5, "f", "=A1+B1"),)
            k3 = (mod.cell_key(6, "f", "=A1+B1"),)
            k4 = (mod.cell_key(5, "f", "=A1+B2"),)
            rows = [{"excel_row": 2, "cells": [], "keys": k1},
                    {"excel_row": 3, "cells": [], "keys": k2},
                    {"excel_row": 4, "cells": [], "keys": k3},
                    {"excel_row": 5, "cells": [], "keys": k4},
                    {"excel_row": 6, "cells": [], "keys": (("formula-missing",),)}]
            groups, skipped = mod.find_exact_groups(rows)
            if groups == {k1: [2, 3]} and skipped == [(6, "公式无缓存值，视为缺失")]:
                ok("4.公式分组单元：字符串+缓存值同时相同才进组；无缓存跳过并给原因")
            else: bad("4.公式分组", "groups=%s skipped=%s" % (groups, skipped))

            # 4b. 含公式的工作簿 → 直接拒绝（端到端；不再宣称可安全输出）
            d = fresh("t4b")
            wb = Workbook(); ws = wb.active; ws.append(HEADER)
            ws.append(["F01", "甲", "X", 1, 2, "2026-01-01"]); ws["D2"] = "=B2*1"
            ws.append(["F01", "甲", "X", 1, 2, "2026-01-01"]); ws["D3"] = "=B2*1"
            wb.save(inp(d))
            h0 = md5(inp(d)); before = tree(d)
            rc, out = run(d)
            after = tree(d)
            leftovers = glob.glob(os.path.join(d, "输出", ".clean-*"))
            if rc == 2 and "公式" in out and md5(inp(d)) == h0 and before == after and not leftovers:
                ok("4b.含公式工作簿拒绝：原始不变、无输出、无残留")
            else: bad("4b.公式拒绝", "rc=%d out=%s" % (rc, out[:150]))

            # 5. 空单元格等同
            d = fresh("t5")
            write_xlsx(inp(d), [
                ["E01", None, "X", 1, 2, "2026-01-01"],
                ["E01", None, "X", 1, 2, "2026-01-01"],
            ])
            rc, out = run(d)
            if "重复组 G1" in out: ok("5.空单元格等同判定")
            else: bad("5.空单元格", out[:150])

            # 6. 内容一致但可能是独立订单：默认保留
            d = fresh("t6")
            write_xlsx(inp(d), [
                ["B001", "陈一", "牛奶", 2, 50, "2026-10-10"],
                ["B001", "陈一", "牛奶", 2, 50, "2026-10-10"],
            ])
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            if ws.max_row == 3 and "标记未删除" in out:
                ok("6.疑似独立订单：默认两条都保留")
            else: bad("6.独立订单", "rows=%d" % ws.max_row)

            # 7. 同订单号不同内容：保留+警告
            d = fresh("t7"); gen(d)
            rc, out = run(d)
            if "疑似重复" in out and "A002" in out:
                ok("7.同订单号不同内容：保留并警告")
            else: bad("7.疑似重复", out[:150])

            # 8. 空白行与纯空格行删除
            d = fresh("t8")
            write_xlsx(inp(d), [
                ["W01", "甲", "X", 1, 2, "2026-01-01"],
                [None, None, None, None, None, None],
                ["   ", "   ", "   ", "   ", "   ", "   "],
            ])
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            if ws.max_row == 2 and stats(out)["删除行数"] == 2:
                ok("8.空白行/纯空格行删除")
            else: bad("8.空白行", "rows=%d" % ws.max_row)

            # 9-17. 清洗规则（用样例断言）
            d = fresh("t9"); gen(d)
            run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            vals = {(ws.cell(row=r, column=1).value): [ws.cell(row=r, column=c).value for c in range(1, 7)]
                    for r in range(2, ws.max_row + 1)}
            checks = [
                ("9.strip", vals["A005"][1] == "钱七"),
                ("10.真实日期格", vals["A001"][5] == "2026-10-01"),
                ("11.日期文本", vals["A002"][5] == "2026-10-02"),
                ("12.模糊日期保留", vals["A007"][5] == "10/07"),
                ("13.金额符号", vals["A003"][4] == 450),
                ("14.金额待定保留", vals["A006"][4] == "待定"),
                ("16.数量单位", vals["A004"][3] == 3),
                ("17.前导零", vals["00012"][0] == "00012"),
            ]
            d2 = fresh("t9b")
            write_xlsx(os.path.join(d2, "输入", "脏订单.xlsx"), [["Q01", "甲", "X", 1, "3,600", "2026-01-01"]])
            run(d2)
            ws2 = load_workbook(os.path.join(d2, "输出", "清理后.xlsx")).active
            checks.append(("13b.千分位", ws2.cell(row=2, column=5).value == 3600))
            d3 = fresh("t9c")
            write_xlsx(os.path.join(d3, "输入", "脏订单.xlsx"), [["Q02", "甲", "X", 1, "1.234", "2026-01-01"]])
            rc3, out3 = run(d3)
            ws3 = load_workbook(os.path.join(d3, "输出", "清理后.xlsx")).active
            checks.append(("15.歧义小数点保留", ws3.cell(row=2, column=5).value == "1.234" and "歧义" in out3))
            failed = [n for n, c in checks if not c]
            if not failed: ok("9-17.清洗规则（strip/日期/金额/数量/前导零/歧义）")
            else: bad("9-17.清洗规则", "失败: %s" % failed)

            # 18. 输入缺失
            d = fresh("t18")
            rc, out = run(d)
            if rc == 2 and "输入文件不存在" in out: ok("18.输入缺失：exit2+中文提示")
            else: bad("18.输入缺失", "rc=%d" % rc)
            # 输入为符号链接 → 拒绝
            d = fresh("t18b"); gen(d)
            os.rename(inp(d), inp(d) + ".real")
            os.symlink(inp(d) + ".real", inp(d))
            rc, out = run(d)
            if rc == 2 and "符号链接" in out: ok("18b.输入符号链接拒绝")
            else: bad("18b.输入链接", "rc=%d out=%s" % (rc, out[:100]))

            # 19. 表为空
            d = fresh("t19")
            wb = Workbook(); wb.active.append(HEADER); wb.save(inp(d))
            rc, out = run(d)
            if rc == 1 and "没有数据行" in out: ok("19.表为空：exit1")
            else: bad("19.表为空", "rc=%d" % rc)

            # 20. 缺少必要列
            d = fresh("t20")
            wb = Workbook(); ws = wb.active
            ws.append(["订单号", "客户"]); ws.append(["X1", "甲"]); wb.save(inp(d))
            rc, out = run(d)
            if rc == 1 and "缺少必要列" in out and "商品" in out: ok("20.缺少必要列：exit1+列出")
            else: bad("20.缺列", "rc=%d" % rc)

            # 21. 输出已存在
            d = fresh("t21"); gen(d)
            r1, _ = run(d)
            r2, o2 = run(d)
            r3, _ = run(d, "--output", "新名.xlsx")
            r4, _ = run(d, "--force")
            if r1 == 0 and r2 == 1 and "拒绝覆盖" in o2 and r3 == 0 and r4 == 0:
                ok("21.输出保护：拒绝/换名/--force")
            else: bad("21.输出保护", "%s" % [r1, r2, r3, r4])

            # 22. 输入输出同一文件
            d = fresh("t22"); gen(d)
            rc, out = run(d, "--output", inp(d))
            rc2, out2 = run(d, "--output", inp(d), "--force")
            if rc == 2 and rc2 == 2:
                ok("22.输入输出同一文件：拒绝（--force亦拒绝）")
            else: bad("22.同文件", "rc=%d/%d" % (rc, rc2))

            # 23/24. 输出目标本身是符号链接（目录内/目录外）→ 一律拒绝
            d = fresh("t23"); gen(d)
            os.makedirs(os.path.join(d, "输出"))
            open(os.path.join(d, "内文件.txt"), "w").write("x")
            os.symlink(os.path.join(d, "内文件.txt"), os.path.join(d, "输出", "清理后.xlsx"))
            rc, out = run(d)
            r1 = (rc == 2 and "符号链接" in out)
            os.unlink(os.path.join(d, "输出", "清理后.xlsx"))
            os.symlink("/tmp/evil23", os.path.join(d, "输出", "清理后.xlsx"))
            rc, out = run(d)
            r2 = (rc == 2 and "符号链接" in out)
            if r1 and r2: ok("23/24.输出目标是符号链接：一律拒绝（目录内/外）")
            else: bad("23/24.链接目标", "%s %s" % (r1, r2))

            # 25. --output 越界
            d = fresh("t25"); gen(d)
            rc, out = run(d, "--output", "../../evil25.xlsx")
            evil = os.path.realpath(os.path.join(d, "输出", "../../evil25.xlsx"))
            if rc == 2 and "演示目录/输出/" in out and not os.path.exists(evil):
                ok("25.--output越界拒绝且未创建文件")
            else: bad("25.越界", "rc=%d" % rc)

            # 26. --force + 不安全路径仍拒绝
            d = fresh("t26"); gen(d)
            rc, out = run(d, "--output", "/tmp/evil26.xlsx", "--force")
            if rc == 2 and not os.path.exists("/tmp/evil26.xlsx"):
                ok("26.--force不安全路径仍拒绝")
            else: bad("26.--force越界", "rc=%d" % rc)

            # 27. 报告文件保护
            d = fresh("t27"); gen(d)
            os.makedirs(os.path.join(d, "输出"))
            open(os.path.join(d, "输出", "清理后_报告.txt"), "w").write("old report")
            rc, out = run(d)
            r1 = (rc == 1 and "报告文件已存在" in out)
            rc, out = run(d, "--force")
            r2 = (rc == 0 and "old report" not in open(os.path.join(d, "输出", "清理后_报告.txt"), encoding="utf-8").read())
            if r1 and r2: ok("27.报告文件：拒绝覆盖/--force可覆盖")
            else: bad("27.报告保护", "%s %s" % (r1, r2))

            # 28/29. --check 严格零写入
            d = fresh("t28"); gen(d)
            before = tree(d)
            rc, out = run(d, "--check")
            after = tree(d)
            r1 = (rc == 0 and before == after and not os.path.exists(os.path.join(d, "输出")))
            d = fresh("t29"); gen(d)
            os.makedirs(os.path.join(d, "输出"))
            open(os.path.join(d, "输出", "旧文件.txt"), "w").write("keep")
            before = tree(d)
            rc, out = run(d, "--check")
            after = tree(d)
            r2 = (rc == 0 and before == after and "原始数据行数：12" in out)
            if r1 and r2: ok("28/29.--check零写入：文件树快照完全一致")
            else: bad("28/29.--check", "%s %s" % (r1, r2))

            # 30. 原子创建：已存在走 FileExistsError，旧字节一致
            d = fresh("t30"); gen(d)
            os.makedirs(os.path.join(d, "输出"))
            p = os.path.join(d, "输出", "清理后.xlsx")
            open(p, "wb").write(b"OLD BYTES")
            rc, out = run(d)
            if rc == 1 and open(p, "rb").read() == b"OLD BYTES" and "拒绝覆盖" in out:
                ok("30.原子防覆盖：旧文件字节一致")
            else: bad("30.原子", "rc=%d" % rc)

            # 31. publish_file 单元测试
            fd, tmp = tempfile.mkstemp(dir=BASE)
            os.close(fd)
            ok1, msg1 = mod.publish_file(tmp, os.path.join(BASE, "no-such-dir", "f.xlsx"), False)
            open(os.path.join(BASE, "ex.xlsx"), "w").write("old")
            fd2, tmp2 = tempfile.mkstemp(dir=BASE); os.close(fd2)
            ok2, msg2 = mod.publish_file(tmp2, os.path.join(BASE, "ex.xlsx"), False)
            still_old = open(os.path.join(BASE, "ex.xlsx")).read() == "old"
            os.unlink(tmp); os.unlink(tmp2)
            if not ok1 and not ok2 and "拒绝覆盖" in msg2 and still_old:
                ok("31.publish_file单元：父目录缺失/已存在拒绝")
            else: bad("31.publish单元", "%s %s" % (ok1, ok2))

            # 32. 报告发布失败 → 部分成功（--force + 故障注入）
            d = fresh("t32"); gen(d)
            run(d)
            old_report = md5(os.path.join(d, "输出", "清理后_报告.txt"))
            rc, out = run(d, "--force", env={"CLEAN_EXCEL_INJECT_REPORT_FAIL": "1"})
            new_xlsx_ok = True
            try:
                load_workbook(os.path.join(d, "输出", "清理后.xlsx")).close()
            except Exception:
                new_xlsx_ok = False
            leftovers = glob.glob(os.path.join(d, "输出", ".clean-*"))
            if rc == 1 and "部分成功" in out and "Excel 已生成" in out and "报告写入失败" in out \
                    and "不是整体成功" in out and md5(os.path.join(d, "输出", "清理后_报告.txt")) == old_report \
                    and new_xlsx_ok and not leftovers:
                ok("32.部分成功：准确报告，不报整体成功，不回滚")
            else: bad("32.部分成功", "rc=%d out=%s" % (rc, out[-300:]))

            # 33. 临时文件无残留
            d = fresh("t33"); gen(d)
            run(d)
            run(d, "--force", env={"CLEAN_EXCEL_INJECT_REPORT_FAIL": "1"})
            leftovers = glob.glob(os.path.join(d, "输出", ".clean-*"))
            if not leftovers: ok("33.临时文件：成功/失败后均无残留")
            else: bad("33.临时残留", str(leftovers))

            # 34. 原始 checksum
            d = fresh("t34"); gen(d)
            h0 = md5(inp(d))
            run(d); run(d, "--force")
            if md5(inp(d)) == h0: ok("34.原始文件checksum不变")
            else: bad("34.checksum", "changed")

            # 35. 教程命令一致性
            tpl = open(os.path.join(ROOT, "data", "tutorials.template.js"), encoding="utf-8").read()
            need = ["--check", "--deduplicate-exact", "--force", "clean_excel.py", "make_dirty_samples.py",
                    "无法可靠地保留公式", ".xlsx"]
            missing = [k for k in need if k not in tpl]
            if not missing: ok("35.教程命令一致性")
            else: bad("35.一致性", "缺: %s" % missing)

            # 36. 哨兵：位于本次测试临时目录内（非主目录固定文件名），且完好
            if sentinel.startswith(sen_dir) \
                    and sentinel != os.path.join(HOME, ".codelab-test-sentinel") \
                    and os.path.isfile(sentinel) \
                    and open(sentinel, encoding="utf-8").read() == "sentinel-ok":
                ok("36.哨兵在测试临时目录内（非主目录固定名），完好")
            else: bad("36.哨兵", sentinel)

            # 37a. 输出目录本身是符号链接 → 指向演示目录自身
            d = fresh("t37a"); gen(d)
            os.symlink(d, os.path.join(d, "输出"))
            before = tree(d)
            rc, out = run(d)
            after = tree(d)
            r_a = (rc == 2 and "符号链接" in out and before == after)

            # 37b. 输出目录 → 演示目录内其他普通子目录
            d = fresh("t37b"); gen(d)
            os.makedirs(os.path.join(d, "备份"))
            os.symlink(os.path.join(d, "备份"), os.path.join(d, "输出"))
            before = tree(d)
            rc, out = run(d)
            after = tree(d)
            r_b = (rc == 2 and "符号链接" in out and before == after)

            # 37c. 输出目录 → 输入目录
            d = fresh("t37c"); gen(d)
            os.symlink(os.path.join(d, "输入"), os.path.join(d, "输出"))
            in_before = tree(os.path.join(d, "输入"))
            rc, out = run(d)
            in_after = tree(os.path.join(d, "输入"))
            r_c = (rc == 2 and "符号链接" in out and in_before == in_after
                   and not os.path.exists(os.path.join(d, "输入", "清理后.xlsx")))

            # 37d. 输出目录 → 演示目录外部
            d = fresh("t37d"); gen(d)
            ext = os.path.join(BASE, "external37d")
            os.makedirs(ext)
            os.symlink(ext, os.path.join(d, "输出"))
            before = tree(d)
            rc, out = run(d)
            after = tree(d)
            r_d = (rc == 2 and "符号链接" in out and before == after and os.listdir(ext) == [])

            # 37e. 输入目录本身是符号链接
            d = fresh("t37e"); gen(d)
            os.rename(os.path.join(d, "输入"), os.path.join(d, "输入.real"))
            os.symlink(os.path.join(d, "输入.real"), os.path.join(d, "输入"))
            before = tree(d)
            rc, out = run(d)
            after = tree(d)
            r_e = (rc == 2 and "符号链接" in out and before == after
                   and not os.path.exists(os.path.join(d, "输出")))
            if r_a and r_b and r_c and r_d and r_e:
                ok("37a-e.输入/输出目录符号链接：一律拒绝，无文件创建/修改")
            else: bad("37.目录链接", "%s %s %s %s %s" % (r_a, r_b, r_c, r_d, r_e))

            # 38. 只有公式、无缓存值的行：不被当空白行删（单元）+ 整表被拒绝（端到端）
            r_blank = mod.is_blank([(None, "f", "=A1")])
            d = fresh("t38")
            wb = Workbook(); ws = wb.active; ws.append(HEADER)
            ws.append([None, None, None, None, None, None]); ws["A2"] = "=B2*0"
            wb.save(inp(d))
            h0 = md5(inp(d)); before = tree(d)
            rc, out = run(d)
            after = tree(d)
            if r_blank is False and rc == 2 and "公式" in out and md5(inp(d)) == h0 and before == after:
                ok("38.公式行不判空白（单元）；含公式表拒绝（端到端），原始不变")
            else: bad("38.公式行", "blank=%s rc=%d" % (r_blank, rc))

            # 40. 部分空白行保留
            d = fresh("t40")
            write_xlsx(inp(d), [["P01", None, "X", 1, 2, "2026-01-01"]])
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            if ws.max_row == 2: ok("40.部分空白行保留")
            else: bad("40.部分空白", "rows=%d" % ws.max_row)

            # 41. 全干净文件修正数为 0
            d = fresh("t41")
            write_xlsx(inp(d), [["C01", "甲", "X", 1, 2, "2026-01-01"]])
            rc, out = run(d)
            if stats(out).get("自动修正单元格数") == 0: ok("41.已干净格不计入修正数")
            else: bad("41.修正数", str(stats(out)))

            # 42. 生成器默认不覆盖已存在文件（字节不变）
            d = fresh("t42"); gen(d)
            p = inp(d)
            open(p, "wb").write(b"USER REAL DATA")
            rc, out = run_gen(d)
            if rc == 1 and "已存在" in out and open(p, "rb").read() == b"USER REAL DATA":
                ok("42.生成器默认不覆盖：用户文件字节不变")
            else: bad("42.生成器覆盖", "rc=%d" % rc)

            # 43. 生成器目标是符号链接 → 拒绝，指向的文件不被修改
            d = fresh("t43"); gen(d)
            p = inp(d)
            os.rename(p, p + ".real")
            open(p + ".real", "wb").write(b"LINK TARGET DATA")
            os.symlink(p + ".real", p)
            rc, out = run_gen(d)
            if rc == 2 and "符号链接" in out and open(p + ".real", "rb").read() == b"LINK TARGET DATA":
                ok("43.生成器符号链接拒绝：目标文件不被修改")
            else: bad("43.生成器链接", "rc=%d out=%s" % (rc, out[:100]))

            # 44. 生成器 --force 允许覆盖
            d = fresh("t44"); gen(d)
            open(inp(d), "wb").write(b"OLD")
            rc, out = run_gen(d, "--force")
            ws = load_workbook(inp(d)).active
            if rc == 0 and ws.max_row == 13:
                ok("44.生成器--force：显式覆盖")
            else: bad("44.生成器force", "rc=%d" % rc)

            # 45. 带时间部分的日期时间 → 拒绝（原始不变、无输出、无残留）
            d = fresh("t45")
            wb = Workbook(); ws = wb.active; ws.append(HEADER)
            ws.append(["D01", "甲", "X", 1, 2, datetime(2026, 1, 1, 10, 30, 0)])
            wb.save(inp(d))
            h0 = md5(inp(d)); before = tree(d)
            rc, out = run(d)
            after = tree(d)
            leftovers = glob.glob(os.path.join(d, "输出", ".clean-*"))
            if rc == 2 and "时间" in out and md5(inp(d)) == h0 and before == after and not leftovers:
                ok("45.带时间日期拒绝：不静默截断，原始不变、无输出、无残留")
            else: bad("45.时间拒绝", "rc=%d out=%s" % (rc, out[:150]))

            # 46. 金额严格规则
            d = fresh("t46")
            write_xlsx(inp(d), [
                ["M01", "甲", "X", 1, "1,2", "2026-01-01"],
                ["M02", "甲", "X", 1, "NaN", "2026-01-01"],
                ["M03", "甲", "X", 1, "Infinity", "2026-01-01"],
                ["M04", "甲", "X", 1, "1,234,567.89", "2026-01-01"],
                ["M05", "甲", "X", 1, "12.5", "2026-01-01"],
            ])
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            vals = [ws.cell(row=r, column=5).value for r in range(2, 7)]
            if rc == 0 and vals == ["1,2", "NaN", "Infinity", 1234567.89, 12.5] \
                    and "非有限" in out and "无法确认" in out and stats(out)["自动修正单元格数"] == 2:
                ok("46.金额严格：1,2/NaN/Infinity保留+警告，合法千分位/小数转换")
            else: bad("46.金额严格", "vals=%s" % vals)

            # 47. 数量严格规则
            d = fresh("t47")
            write_xlsx(inp(d), [
                ["Q01", "甲", "X", "1件2", 100, "2026-01-01"],
                ["Q02", "甲", "X", "3.5", 100, "2026-01-01"],
                ["Q03", "甲", "X", "4箱", 100, "2026-01-01"],
                ["Q04", "甲", "X", 6, 100, "2026-01-01"],
                ["Q05", "甲", "X", "NaN", 100, "2026-01-01"],
            ])
            rc, out = run(d)
            ws = load_workbook(os.path.join(d, "输出", "清理后.xlsx")).active
            vals = [ws.cell(row=r, column=4).value for r in range(2, 7)]
            if rc == 0 and vals == ["1件2", "3.5", 4, 6, "NaN"] \
                    and "不是整数" in out and "无法确认" in out and "非有限" in out:
                ok("47.数量严格：1件2/3.5/NaN保留+警告，合法单位转换")
            else: bad("47.数量严格", "vals=%s" % vals)

            # 48. 类型/时间/错误值（内部函数单元测试）
            r1 = {"excel_row": 2, "cells": [], "keys": (mod.scalar_key(datetime(2026, 1, 1, 10, 0)),)}
            r2 = {"excel_row": 3, "cells": [], "keys": (mod.scalar_key(datetime(2026, 1, 1, 11, 0)),)}
            g1, _ = mod.find_exact_groups([r1, r2])
            c1 = (g1 == {})
            r3 = {"excel_row": 4, "cells": [], "keys": (mod.scalar_key(10),)}
            r4 = {"excel_row": 5, "cells": [], "keys": (mod.scalar_key(10.0),)}
            g2, _ = mod.find_exact_groups([r3, r4])
            c2 = (len(g2) == 1)  # int 10 与 float 10.0 视为相同（已在文档中声明）
            r5 = {"excel_row": 6, "cells": [], "keys": (("error",),)}
            r6 = {"excel_row": 7, "cells": [], "keys": (("error",),)}
            g3, s3 = mod.find_exact_groups([r5, r6])
            c3 = (g3 == {} and len(s3) == 2 and "错误值" in s3[0][1])
            r7 = {"excel_row": 8, "cells": [], "keys": (mod.scalar_key(float("nan")),)}
            g4, s4 = mod.find_exact_groups([r7, {"excel_row": 9, "cells": [], "keys": (mod.scalar_key(float("inf")),)}])
            c4 = (g4 == {} and len(s4) == 2)
            if c1 and c2 and c3 and c4:
                ok("48.类型规则：不同时间不进组；int/float相同；错误值/非有限跳过+原因")
            else: bad("48.类型规则", "%s %s %s %s" % (c1, c2, c3, c4))

            # 49. --output 扩展名拒绝
            d = fresh("t49"); gen(d)
            before = tree(d)
            rc, out = run(d, "--output", "x.csv")
            rc2, out2 = run(d, "--output", "x.txt", "--force")
            after = tree(d)
            if rc == 2 and rc2 == 2 and ".xlsx" in out and before == after:
                ok("49.--output扩展名：非.xlsx拒绝（--force亦拒绝），无文件创建")
            else: bad("49.扩展名", "rc=%d/%d" % (rc, rc2))

        print("\n==== 共 %d 通过, %d 失败 ====" % (PASS, FAIL))
        return 1 if FAIL else 0
    finally:
        # 只清理本次测试创建的哨兵临时目录
        shutil.rmtree(sen_dir, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
